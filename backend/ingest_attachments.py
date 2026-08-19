"""ingest_attachments.py — broker email (.eml/.msg) -> Submission TEXT, fully local.

Scope (per the build brief): read enough TEXT from a broker email and its
attachments to (a) classify the whole submission into buckets and (b) index each
attachment by document type. NO field extraction, NO OCR, NO cloud calls.

- .eml  : Python stdlib `email` (nested message/rfc822 forwards are unwrapped)
- .msg  : python-oxmsg (MIT), best-effort — the demo set is .eml
- PDF   : pdfplumber        DOCX: python-docx        XLSX: openpyxl       CSV: as-is
- Anything unreadable sets `read_warning` and continues — never crash the demo.
- Deterministic: same input file -> byte-identical output, so scoring stays honest.
- PII: obvious identifiers (emails, phones, SSN/FEIN, street addresses) are replaced
  with [PLACEHOLDER_n] BEFORE the text can reach any LLM. Pattern-based, best-effort:
  person names without a pattern are NOT caught — do not treat this as a redaction
  guarantee. Company names are deliberately kept (they are classification signal).
"""
import io, os, re, json, email, hashlib
from email import policy
from html.parser import HTMLParser
from pydantic import BaseModel

# per-attachment text cap: enough for doc-type classification, keeps SOV workbooks
# with thousands of rows from swamping the submission text
MAX_ATT_CHARS = 30_000
LOW_TEXT_CHARS = 200          # a "real" document under this is probably a scan

TEXT_EXT = {".txt", ".csv"}
PDF_EXT = {".pdf"}
DOCX_EXT = {".docx"}
XLSX_EXT = {".xlsx", ".xlsm"}
SKIP_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}   # signature logos etc.; no OCR


class Attachment(BaseModel):
    filename: str
    doc_type: str | None = None      # filled by the classifier channel, not here
    text: str
    char_count: int
    read_warning: str | None = None


class Submission(BaseModel):
    submission_id: str
    cover_email_text: str
    attachments: list[Attachment]
    combined_text: str               # cover email + all attachment text


# ---------- PII anonymiser (pattern-based, deterministic) ----------
# Order matters: more specific patterns first so e.g. an SSN inside an address
# line is tokenised as itself. All matches share one placeholder table per
# submission, assigned in first-appearance order (deterministic).
_PII_PATTERNS = [
    re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"),                  # email
    re.compile(r"\b\d{3}-\d{2}-\d{4}\b"),                                            # SSN
    re.compile(r"\bFEIN[:\s#]*\d{2}-?\d{7}\b", re.I),                                # FEIN (labelled)
    re.compile(r"\b\d{2}-\d{7}\b"),                                                  # FEIN (bare)
    re.compile(r"(?:\+1[\s.-]?)?\(?\b\d{3}\)?[\s.-]\d{3}[\s.-]\d{4}\b"),             # US phone
    re.compile(r"\b\d{1,5}\s+[A-Z][A-Za-z]*(?:\s+[A-Z][A-Za-z]*){0,3}\s+"
               r"(?:Street|St|Avenue|Ave|Road|Rd|Drive|Dr|Boulevard|Blvd|Lane|Ln|"
               r"Way|Court|Ct|Parkway|Pkwy|Highway|Hwy|Place|Pl|Circle|Cir)\.?\b"),  # street address
]


# display names on mail-header lines ("From: Jane Doe <addr>") — the one place
# person names sit in a reliable pattern. Names elsewhere in prose are NOT caught.
_HEADER_NAME = re.compile(r"(?m)^((?:From|To|Cc|Bcc)\s*:\s*)([^<\n]{2,80}?)(\s*[<\[])")


def anonymize(text, table):
    """Replace PII matches with [PLACEHOLDER_n]. `table` (match->placeholder) is
    shared across a whole submission so the same address gets the same token."""
    def tok(s):
        if s not in table:
            table[s] = f"[PLACEHOLDER_{len(table) + 1}]"
        return table[s]
    text = _HEADER_NAME.sub(lambda m: m.group(1) + tok(m.group(2).strip()) + m.group(3), text)
    for pat in _PII_PATTERNS:
        text = pat.sub(lambda m: tok(m.group(0)), text)
    return text


# ---------- attachment bytes -> text ----------
def _pdf_text(data):
    import pdfplumber
    out = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            out.append(page.extract_text() or "")
    return "\n".join(out)


def _docx_text(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(p for p in parts if p.strip())


def _xlsx_text(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[sheet: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append("\t".join(cells))
    wb.close()
    return "\n".join(out)


def extract_attachment(filename, data):
    """One attachment -> Attachment (text + optional warning). Never raises."""
    ext = os.path.splitext(filename)[1].lower()
    text, warning = "", None
    try:
        if ext in SKIP_EXT:
            warning = "image attachment — skipped (no OCR)"
        elif ext in TEXT_EXT:
            text = data.decode("utf-8", errors="replace")
        elif ext in PDF_EXT:
            text = _pdf_text(data)
        elif ext in DOCX_EXT:
            text = _docx_text(data)
        elif ext in XLSX_EXT:
            text = _xlsx_text(data)
        else:
            warning = f"unsupported type ({ext or 'no extension'}) — skipped"
    except Exception as e:
        warning = f"read failed ({type(e).__name__}) — skipped"
        text = ""
    text = text.strip()
    if not warning and len(text) < LOW_TEXT_CHARS:
        warning = "scanned or low text yield"
    # cap silently — hitting the cap is normal for big ACORD packets/SOVs, and a
    # deterministic cap is not a read problem worth a ⚠ in the dashboard
    if len(text) > MAX_ATT_CHARS:
        text = text[:MAX_ATT_CHARS]
    return Attachment(filename=filename, text=text, char_count=len(text),
                      read_warning=warning)


# ---------- email body -> text ----------
class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(); self.parts = []; self._skip = 0
    def handle_starttag(self, tag, attrs):
        if tag in ("style", "script"): self._skip += 1
    def handle_endtag(self, tag):
        if tag in ("style", "script"): self._skip = max(0, self._skip - 1)
        if tag in ("p", "div", "br", "tr", "li"): self.parts.append("\n")
    def handle_data(self, d):
        if not self._skip: self.parts.append(d)


def _html_to_text(html):
    p = _HTMLText(); p.feed(html)
    return re.sub(r"\n{3,}", "\n\n", "".join(p.parts))


# ---------- .eml parsing (stdlib; unwraps forwarded message/rfc822) ----------
def _walk_eml(msg, plain, html, atts):
    for part in msg.walk():
        ctype = part.get_content_type()
        fn = part.get_filename()
        if ctype == "message/rfc822":
            continue                       # walk() already descends into it
        if fn:
            atts.append((fn, part.get_payload(decode=True) or b""))
        elif ctype == "text/plain":
            try: plain.append(part.get_content())
            except Exception: pass
        elif ctype == "text/html":
            try: html.append(part.get_content())
            except Exception: pass


def parse_eml(path):
    """-> (subject, body_text, [(filename, bytes), ...])"""
    with open(path, "rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)
    plain, html, atts = [], [], []
    _walk_eml(msg, plain, html, atts)
    body = "\n".join(plain).strip() or _html_to_text("\n".join(html)).strip()
    return (msg["subject"] or ""), body, atts


# ---------- .msg parsing (best-effort; the demo set is .eml) ----------
def parse_msg(path):
    from oxmsg import Message
    m = Message.load(path)
    body = (getattr(m, "body", None) or "").strip()
    if not body:
        body = _html_to_text(getattr(m, "html_body", "") or "").strip()
    atts = [(a.file_name or "attachment", a.file_bytes or b"")
            for a in getattr(m, "attachments", [])]
    return (getattr(m, "subject", "") or ""), body, atts


# ---------- the entry point ----------
def submission_id_for(path):
    """Deterministic slug of the filename (matched against ground truth the same
    way ingest.py matches: case/punctuation-insensitive)."""
    stem = os.path.splitext(os.path.basename(path))[0]
    stem = re.sub(r"^(fw|fwd|re)[_\s:]+", "", stem, flags=re.I)
    return re.sub(r"[^a-z0-9]+", "_", stem.lower()).strip("_")


def load_submission(path, submission_id=None):
    """One .eml/.msg file -> anonymised Submission. Deterministic; never raises
    on unreadable attachments (they carry read_warning instead)."""
    ext = os.path.splitext(path)[1].lower()
    subject, body, raw_atts = (parse_msg if ext == ".msg" else parse_eml)(path)
    table = {}                                   # one PII table per submission
    cover = anonymize(f"SUBJECT: {subject}\n\n{body}".strip(), table)
    attachments = []
    for fn, data in raw_atts:
        a = extract_attachment(fn, data)
        a.text = anonymize(a.text, table)
        a.char_count = len(a.text)
        attachments.append(a)
    combined = "\n\n".join(
        ["=== COVER EMAIL ===", cover]
        + [x for a in attachments for x in (f"=== ATTACHMENT: {a.filename} ===", a.text)])
    return Submission(submission_id=submission_id or submission_id_for(path),
                      cover_email_text=cover, attachments=attachments,
                      combined_text=combined)


def content_hash(sub):
    """Stable hash of everything the classifier would see — the determinism check."""
    return hashlib.sha256(json.dumps(sub.model_dump(), sort_keys=True)
                          .encode()).hexdigest()[:16]
